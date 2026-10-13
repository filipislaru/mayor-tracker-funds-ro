"""Tests for normalising local budget execution datasets with synthetic fixtures."""

from __future__ import annotations

import datetime as dt
import json
from decimal import Decimal
from pathlib import Path

import openpyxl
import pandas as pd
import pyarrow as pa
import pyarrow.parquet as pq
import pytest
from mayortracker.normalise.budget_exec import (
    _col_idx_from_letter,
    _to_decimal,
    load_budget_columns_config,
    normalise_budget_exec,
    parse_annual_budget_file,
)
from mayortracker.util.config import Settings, find_repo_root


def test_col_idx_from_letter() -> None:
    assert _col_idx_from_letter("A") == 0
    assert _col_idx_from_letter("B") == 1
    assert _col_idx_from_letter("U") == 20
    assert _col_idx_from_letter("V") == 21
    assert _col_idx_from_letter("W") == 22
    assert _col_idx_from_letter("Z") == 25
    assert _col_idx_from_letter("AA") == 26
    assert _col_idx_from_letter("AF") == 31
    assert _col_idx_from_letter("AH") == 33
    assert _col_idx_from_letter("AJ") == 35
    assert _col_idx_from_letter("AK") == 36


def test_to_decimal() -> None:
    assert _to_decimal(None) == Decimal("0.00")
    assert _to_decimal("") == Decimal("0.00")
    assert _to_decimal(100) == Decimal("100.00")
    assert _to_decimal(123.456) == Decimal("123.46")
    assert _to_decimal("1,234.50") == Decimal("1234.50")
    assert _to_decimal("invalid") == Decimal("0.00")
    assert _to_decimal(Decimal("55.10")) == Decimal("55.10")


def _create_synthetic_excel(file_path: Path, mismatch_code: bool = False) -> None:
    """Create a synthetic Anexa 24 workbook obeying the 2024 column layout."""
    wb = openpyxl.Workbook()
    # Create sheets: summary sheet and county sheet "1"
    ws_total = wb.active
    ws_total.title = "total TARA"
    ws = wb.create_sheet(title="1")

    # Fill rows up to row 13
    for _r in range(1, 14):
        ws.append([None] * 40)

    # Row 2 (index 2): County header
    ws.cell(row=2, column=2, value="Judeţul: TEST ALBA")

    # Row 9-11: Labels
    ws.cell(
        row=9,
        column=21,
        value="Sume primite de la UE/alti donatori in contul platilor efectuate si prefinantari",
    )  # U
    ws.cell(row=9, column=22, value="Alte sume primite de la UE")  # V
    ws.cell(
        row=9,
        column=23,
        value=(
            "Sume primite de la UE/alţi donatori în contul plăţilor "
            "efectuate şi prefinanţării aferente cadrului financiar 2014-2020"
        ),
    )  # W

    ws.cell(
        row=11, column=32, value="Proiecte cu finanţare din fonduri externe nerambursabile"
    )  # AF
    ws.cell(
        row=11,
        column=34,
        value=(
            "Proiecte cu finanţare din fonduri externe nerambursabile "
            "aferente cadrului financiar 2014-2020"
        ),
    )  # AH
    ws.cell(
        row=11,
        column=36,
        value=(
            "Proiecte cu finantare din sumele reprezentând "
            "asistenţa financiară nerambursabilă aferentă PNRR"
        ),
    )  # AJ
    ws.cell(
        row=11,
        column=37,
        value="Proiecte cu finantare din sumele aferente componentei de împrumut a PNRR",
    )  # AK

    # Row 12: Codes
    ws.cell(row=12, column=21, value="45.02")
    ws.cell(row=12, column=22, value="46.02")
    ws.cell(row=12, column=23, value="99.99" if mismatch_code else "48.02")
    ws.cell(row=12, column=32, value="56")
    ws.cell(row=12, column=34, value="58")
    ws.cell(row=12, column=36, value="60")
    ws.cell(row=12, column=37, value="61")

    # Add locality rows:
    # 1. Standard UAT: 900001
    row_1 = [None] * 40
    row_1[0] = 1  # Line number in col A
    row_1[1] = 900001  # SIRUTA in col B
    row_1[2] = "TEST COMUNA UNU"  # Name in col C
    row_1[20] = 200000  # Col U (45.02)
    row_1[21] = 10000  # Col V (46.02)
    row_1[22] = 1000000  # Col W (48.02)
    row_1[31] = 50000  # Col AF (56)
    row_1[33] = 900000  # Col AH (58)
    row_1[35] = 25000  # Col AJ (60)
    row_1[36] = 10000  # Col AK (61)
    ws.append(row_1)

    # 2. Locality mapping row: raw SIRUTA 900099 (village under parent UAT 900002)
    row_2 = [None] * 40
    row_2[0] = 2
    row_2[1] = 900099
    row_2[2] = "TEST SAT COMPONENT"
    row_2[22] = 500000  # 48.02
    row_2[33] = 400000  # 58
    ws.append(row_2)

    # 3. Duplicate row for UAT 900003: First an empty placeholder
    row_3a = [None] * 40
    row_3a[0] = 3
    row_3a[1] = 900003
    row_3a[2] = "TEST COMUNA TREI"
    ws.append(row_3a)

    # 4. Duplicate row for UAT 900003: Populated active row
    row_3b = [None] * 40
    row_3b[0] = 4
    row_3b[1] = 900003
    row_3b[2] = "TEST COMUNA TREI"
    row_3b[22] = 300000  # 48.02
    row_3b[33] = 150000  # 58
    ws.append(row_3b)

    # 5. Outlier ratio row: UAT 900004 (spending >> revenue)
    row_4 = [None] * 40
    row_4[0] = 5
    row_4[1] = 900004
    row_4[2] = "TEST ORAS PATRU"
    row_4[22] = 1000  # 48.02
    row_4[33] = 500000  # 58 (ratio = 500.0, outlier)
    ws.append(row_4)

    # 6. Unmapped village row without override: raw SIRUTA 900088
    row_5 = [None] * 40
    row_5[0] = 6
    row_5[1] = 900088
    row_5[2] = "TEST SAT FARA OVERRIDE"
    row_5[22] = 20000  # 48.02
    row_5[33] = 20000  # 58 (ratio = 1.0, normal)
    ws.append(row_5)

    wb.save(file_path)


def test_parse_annual_budget_file_layout_mismatch(tmp_path: Path) -> None:
    excel_path = tmp_path / "mismatch.xlsx"
    _create_synthetic_excel(excel_path, mismatch_code=True)
    repo_root = find_repo_root()
    config = load_budget_columns_config(repo_root)

    uat_ref = pd.DataFrame(
        [
            {
                "siruta": 900001,
                "name": "TEST COMUNA UNU",
                "county_name": "TEST ALBA",
                "county_code": "AB",
            }
        ]
    )
    provenance = {
        "source_id": "budget_exec",
        "source_url": "https://example.com/test",
        "retrieved_at": dt.datetime.now(dt.UTC),
        "raw_sha256": "fakehash",
    }

    with pytest.raises(ValueError, match="Column layout mismatch"):
        parse_annual_budget_file(
            file_path=excel_path,
            year=2024,
            year_config=config["2024"],
            uat_ref=uat_ref,
            siruta_overrides={},
            provenance=provenance,
        )


def test_normalise_budget_exec_synthetic_pipeline(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    repo_root = find_repo_root()
    data_dir = tmp_path / "data"
    raw_dir = data_dir / "raw" / "budget_exec" / "2026-10-13"
    raw_dir.mkdir(parents=True, exist_ok=True)
    processed_dir = data_dir / "processed"
    processed_dir.mkdir(parents=True, exist_ok=True)

    # Create synthetic uat.parquet
    uat_schema = pa.schema(
        [
            ("siruta", pa.int64()),
            ("name", pa.string()),
            ("county_name", pa.string()),
            ("county_code", pa.string()),
        ]
    )
    uat_data = {
        "siruta": [900001, 900002, 900003, 900004],
        "name": ["TEST COMUNA UNU", "TEST ORAS DOI", "TEST COMUNA TREI", "TEST ORAS PATRU"],
        "county_name": ["TEST ALBA", "TEST ALBA", "TEST ALBA", "TEST ALBA"],
        "county_code": ["AB", "AB", "AB", "AB"],
    }
    pq.write_table(pa.Table.from_pydict(uat_data, schema=uat_schema), processed_dir / "uat.parquet")

    # Mock siruta overrides to map village 900099 -> UAT 900002 (900088 is unmapped)
    monkeypatch.setattr(
        "mayortracker.normalise.budget_exec.load_siruta_overrides",
        lambda _root: {900099: 900002},
    )

    # Create synthetic Excel file and manifest
    excel_path = raw_dir / "situaia-veniturilor-i-cheltuielilor-2024.xlsx"
    _create_synthetic_excel(excel_path, mismatch_code=False)

    manifest_entries = [
        {
            "filename": "situaia-veniturilor-i-cheltuielilor-2024.xlsx",
            "url": "https://example.com/test-2024.xlsx",
            "retrieved_at": "2026-10-13T12:00:00+00:00",
            "sha256": "fakehash123",
            "bytes": 1024,
            "http_status": 200,
        }
    ]
    (raw_dir / "manifest.json").write_text(json.dumps(manifest_entries), encoding="utf-8")

    settings = Settings(
        project_name="TEST project",
        data_dir=data_dir,
        contact_email="test@example.com",
        config_path=repo_root / "config" / "settings.yaml",
    )

    metrics = normalise_budget_exec(settings)

    assert metrics.years_processed == [2024]
    # 4 UATs + 1 unmatched row = 5 rows produced
    assert metrics.total_uats_produced == 5
    assert metrics.matched_uats_count == 4
    assert metrics.unmatched_uats_count == 1
    assert metrics.unmatched_details[0]["raw_siruta"] == 900088

    # Verify uat_budget_eu_lines.parquet
    lines_path = processed_dir / "uat_budget_eu_lines.parquet"
    assert lines_path.is_file()
    lines_table = pq.read_table(lines_path)
    # 5 produced * 7 lines = 35 lines
    assert len(lines_table) == 35
    assert lines_table.schema.field("amount").type == pa.decimal128(20, 2)

    # Verify uat_eu_received_year.parquet
    summary_path = processed_dir / "uat_eu_received_year.parquet"
    assert summary_path.is_file()
    summary_table = pq.read_table(summary_path)
    assert len(summary_table) == 5
    df = summary_table.to_pandas()

    # Verify UAT 1 values
    u1 = df[df["siruta"] == 900001].iloc[0]
    assert u1["eu_received_2014_2020"] == Decimal("1000000.00")
    assert u1["eu_received_unspecified"] == Decimal("200000.00")
    assert u1["eu_received_other"] == Decimal("10000.00")
    assert u1["eu_received_total"] == Decimal("1210000.00")
    assert u1["eu_spent_2014_2020"] == Decimal("900000.00")
    assert u1["eu_spent_unspecified"] == Decimal("50000.00")
    assert u1["eu_spent_pnrr_grant"] == Decimal("25000.00")
    assert u1["eu_spent_pnrr_loan"] == Decimal("10000.00")
    assert u1["eu_spent_total"] == Decimal("985000.00")
    assert round(u1["ratio_2014_2020"], 2) == 0.90

    # Verify village 900099 was mapped to UAT 2 (900002) via override
    u2 = df[df["siruta"] == 900002].iloc[0]
    assert u2["eu_received_2014_2020"] == Decimal("500000.00")
    assert u2["eu_spent_2014_2020"] == Decimal("400000.00")

    # Verify duplicate placeholder handling on UAT 3 (900003)
    u3 = df[df["siruta"] == 900003].iloc[0]
    assert u3["eu_received_2014_2020"] == Decimal("300000.00")
    assert u3["eu_spent_2014_2020"] == Decimal("150000.00")

    # Verify outlier ratio on UAT 4 (900004)
    u4 = df[df["siruta"] == 900004].iloc[0]
    assert u4["ratio_2014_2020"] == 500.0
    assert metrics.outlier_count_by_year[2024] == 1
