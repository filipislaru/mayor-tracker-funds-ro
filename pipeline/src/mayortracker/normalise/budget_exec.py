"""Normalise local budget execution datasets into clean Parquet tables."""

from __future__ import annotations

import datetime as dt
import json
import logging
from decimal import Decimal
from pathlib import Path
from typing import Any

import duckdb
import openpyxl
import pandas as pd
import pyarrow as pa
import pyarrow.parquet as pq
import yaml
from pydantic import BaseModel

from mayortracker.schemas.budget_exec import BudgetEuLineRecord, UatEuReceivedYearRecord
from mayortracker.util.config import Settings, find_repo_root

logger = logging.getLogger(__name__)


class BudgetExecMetrics(BaseModel):
    """Execution metrics from normalising budget execution datasets."""

    years_processed: list[int]
    total_raw_rows: int
    total_uats_produced: int
    matched_uats_count: int
    unmatched_uats_count: int
    unmatched_details: list[dict[str, Any]]
    total_eu_received_by_year: dict[int, Decimal]
    total_eu_spent_by_year: dict[int, Decimal]
    outlier_count_by_year: dict[int, int]


def load_budget_columns_config(repo_root: Path) -> dict[str, Any]:
    """Load per-year column mapping from config/budget_exec_columns.yaml."""
    config_path = repo_root / "config" / "budget_exec_columns.yaml"
    if not config_path.is_file():
        raise FileNotFoundError(f"Missing column configuration file: {config_path}")
    with config_path.open("r", encoding="utf-8") as f:
        data = yaml.safe_load(f) or {}
    return data


def load_siruta_overrides(repo_root: Path) -> dict[int, int]:
    """Load verified SIRUTA code overrides from config/siruta_overrides.yaml."""
    path = repo_root / "config" / "siruta_overrides.yaml"
    if not path.is_file():
        return {}
    with path.open("r", encoding="utf-8") as f:
        data = yaml.safe_load(f) or {}
    overrides_raw = data.get("overrides", {})
    return {int(k): int(v["target_uat_siruta"]) for k, v in overrides_raw.items()}


def _to_decimal(val: Any) -> Decimal:
    """Convert an Excel cell value to Decimal, rounding to 2 decimal places."""
    if val is None or val == "":
        return Decimal("0.00")
    if isinstance(val, (int, Decimal)):
        return Decimal(val).quantize(Decimal("0.01"))
    if isinstance(val, float):
        return Decimal(str(round(val, 2))).quantize(Decimal("0.01"))
    val_str = str(val).strip().replace(",", "")
    try:
        return Decimal(val_str).quantize(Decimal("0.01"))
    except Exception:
        return Decimal("0.00")


def _col_idx_from_letter(col_letter: str) -> int:
    """Convert column letter (e.g. 'A', 'U', 'AK') to 0-based integer index."""
    num = 0
    for char in col_letter.upper():
        num = num * 26 + (ord(char) - ord("A") + 1)
    return num - 1


def parse_annual_budget_file(
    file_path: Path,
    year: int,
    year_config: dict[str, Any],
    uat_ref: pd.DataFrame,
    siruta_overrides: dict[int, int],
    provenance: dict[str, Any],
) -> tuple[list[BudgetEuLineRecord], list[UatEuReceivedYearRecord], dict[str, Any]]:
    """Parse one annual budget execution workbook and produce line and summary records."""
    uat_meta = uat_ref.set_index("siruta").to_dict(orient="index")

    layout = year_config.get("sheet_layout", {})
    code_row_num = layout.get("code_row", 12)
    siruta_col_idx = _col_idx_from_letter(layout.get("siruta_col", "B"))
    uat_name_col_idx = _col_idx_from_letter(layout.get("uat_name_col", "C"))

    revenue_lines_cfg = year_config.get("eu_revenue_lines", [])
    expenditure_lines_cfg = year_config.get("eu_expenditure_lines", [])

    wb = openpyxl.load_workbook(file_path, read_only=True, data_only=True)
    summary_sheets = {"total TARA", "total COMUNE", "total ORASE", "total MUNICIPII", "total CJ"}
    county_sheets = [s for s in wb.sheetnames if s not in summary_sheets]

    # Verify column codes in the first county sheet against config
    if county_sheets:
        first_ws = wb[county_sheets[0]]
        for r_idx, row in enumerate(first_ws.iter_rows(values_only=True), start=1):
            if r_idx == code_row_num:
                for line in revenue_lines_cfg + expenditure_lines_cfg:
                    target_idx = _col_idx_from_letter(line["col_letter"])
                    actual_code = (
                        str(row[target_idx] or "").strip() if target_idx < len(row) else ""
                    )
                    expected_code = str(line["code"]).strip()
                    if actual_code != expected_code:
                        msg = (
                            f"Column layout mismatch in {file_path.name} "
                            f"sheet '{county_sheets[0]}': col {line['col_letter']} "
                            f"expected code '{expected_code}', found '{actual_code}'. "
                            f"Update config/budget_exec_columns.yaml."
                        )
                        raise ValueError(msg)
                break

    all_line_records: list[BudgetEuLineRecord] = []
    summary_records_map: dict[int, UatEuReceivedYearRecord] = {}

    unmatched_entries: list[dict[str, Any]] = []
    total_raw_rows = 0

    # Temporary collector to handle possible duplicate / placeholder rows per UAT
    raw_uat_rows: dict[int, list[dict[str, Any]]] = {}

    for sheet_name in county_sheets:
        ws = wb[sheet_name]
        county_header = ""
        for row in ws.iter_rows(values_only=True):
            if not row or not any(row):
                continue
            if row[1] and str(row[1]).strip().startswith("Judeţul:"):
                county_header = str(row[1]).replace("Judeţul:", "").strip()
                continue

            col_a = str(row[0] or "").strip()
            col_siruta = str(row[siruta_col_idx] or "").strip() if siruta_col_idx < len(row) else ""
            col_name = (
                str(row[uat_name_col_idx] or "").strip() if uat_name_col_idx < len(row) else ""
            )

            # Locality rows have numeric line number in col A and numeric SIRUTA in col B
            if col_a.isdigit() and col_siruta.isdigit():
                raw_siruta = int(col_siruta)
                total_raw_rows += 1

                # Check if SIRUTA has an explicit verified override (e.g. Racsa 138404 -> 180091)
                effective_siruta = siruta_overrides.get(raw_siruta, raw_siruta)

                # Collect row data
                row_vals = {}
                row_vals["_raw_siruta"] = raw_siruta
                row_vals["_effective_siruta"] = effective_siruta
                row_vals["_name"] = col_name
                row_vals["_county"] = county_header

                # Read revenue values
                for r_line in revenue_lines_cfg:
                    idx = _col_idx_from_letter(r_line["col_letter"])
                    val = row[idx] if idx < len(row) else None
                    row_vals[f"rev_{r_line['code']}"] = _to_decimal(val)

                # Read expenditure values
                for e_line in expenditure_lines_cfg:
                    idx = _col_idx_from_letter(e_line["col_letter"])
                    val = row[idx] if idx < len(row) else None
                    row_vals[f"exp_{e_line['code']}"] = _to_decimal(val)

                # Total row amount across all extracted fields
                # (used to disambiguate zero placeholders)
                row_vals["_total_activity"] = sum(
                    [row_vals[f"rev_{r['code']}"] for r in revenue_lines_cfg]
                    + [row_vals[f"exp_{e['code']}"] for e in expenditure_lines_cfg]
                )

                raw_uat_rows.setdefault(effective_siruta, []).append(row_vals)

    # Disambiguate rows (select the active row if duplicate placeholders exist)
    for eff_siruta, rlist in raw_uat_rows.items():
        if len(rlist) == 1:
            chosen = rlist[0]
        else:
            # Sort by total activity descending, then non-empty name
            chosen = sorted(
                rlist, key=lambda x: (x["_total_activity"], len(x["_name"])), reverse=True
            )[0]

        # Check UAT matching
        uat_info = uat_meta.get(eff_siruta)
        if uat_info:
            matched_siruta = eff_siruta
            uat_name = uat_info["name"]
            county_name = uat_info["county_name"]
            county_code = uat_info["county_code"]
        else:
            matched_siruta = eff_siruta
            uat_name = chosen["_name"]
            county_name = chosen["_county"]
            county_code = None
            unmatched_entries.append(
                {
                    "raw_siruta": chosen["_raw_siruta"],
                    "effective_siruta": eff_siruta,
                    "name": chosen["_name"],
                    "county": chosen["_county"],
                }
            )

        # Build BudgetEuLineRecords
        for r_line in revenue_lines_cfg:
            code = r_line["code"]
            amt = chosen[f"rev_{code}"]
            all_line_records.append(
                BudgetEuLineRecord(
                    siruta=matched_siruta,
                    year=year,
                    classification_code=str(code),
                    label=r_line["label"],
                    line_type="revenue",
                    amount=amt,
                    unit="RON",
                    source_id=provenance["source_id"],
                    source_url=provenance["source_url"],
                    retrieved_at=provenance["retrieved_at"],
                    raw_sha256=provenance["raw_sha256"],
                )
            )

        for e_line in expenditure_lines_cfg:
            code = e_line["code"]
            amt = chosen[f"exp_{code}"]
            all_line_records.append(
                BudgetEuLineRecord(
                    siruta=matched_siruta,
                    year=year,
                    classification_code=str(code),
                    label=e_line["label"],
                    line_type="expenditure",
                    amount=amt,
                    unit="RON",
                    source_id=provenance["source_id"],
                    source_url=provenance["source_url"],
                    retrieved_at=provenance["retrieved_at"],
                    raw_sha256=provenance["raw_sha256"],
                )
            )

        # Build UatEuReceivedYearRecord
        rec_2014_2020 = chosen.get("rev_48.02", Decimal("0.00"))
        rec_unspecified = chosen.get("rev_45.02", Decimal("0.00"))
        rec_other = chosen.get("rev_46.02", Decimal("0.00"))
        rec_total = rec_2014_2020 + rec_unspecified + rec_other

        spent_2014_2020 = chosen.get("exp_58", Decimal("0.00"))
        spent_unspecified = chosen.get("exp_56", Decimal("0.00"))
        spent_pnrr_grant = chosen.get("exp_60", Decimal("0.00"))
        spent_pnrr_loan = chosen.get("exp_61", Decimal("0.00"))
        spent_total = spent_2014_2020 + spent_unspecified + spent_pnrr_grant + spent_pnrr_loan

        # Ratios (consistency checks)
        spending_to_revenue_ratio = (
            float(spent_total / rec_total) if rec_total > Decimal("0.00") else None
        )
        ratio_2014_2020 = (
            float(spent_2014_2020 / rec_2014_2020) if rec_2014_2020 > Decimal("0.00") else None
        )

        summary_records_map[matched_siruta] = UatEuReceivedYearRecord(
            siruta=matched_siruta,
            year=year,
            uat_name=uat_name,
            county_name=county_name,
            county_code=county_code,
            eu_received_total=rec_total,
            eu_received_2014_2020=rec_2014_2020,
            eu_received_unspecified=rec_unspecified,
            eu_received_other=rec_other,
            eu_spent_total=spent_total,
            eu_spent_2014_2020=spent_2014_2020,
            eu_spent_unspecified=spent_unspecified,
            eu_spent_pnrr_grant=spent_pnrr_grant,
            eu_spent_pnrr_loan=spent_pnrr_loan,
            spending_to_revenue_ratio=spending_to_revenue_ratio,
            ratio_2014_2020=ratio_2014_2020,
            unit="RON",
            source_id=provenance["source_id"],
            source_url=provenance["source_url"],
            retrieved_at=provenance["retrieved_at"],
            raw_sha256=provenance["raw_sha256"],
        )

    summary_records = list(summary_records_map.values())
    metrics_info = {
        "total_raw_rows": total_raw_rows,
        "uats_produced": len(summary_records),
        "unmatched": unmatched_entries,
    }
    return all_line_records, summary_records, metrics_info


def normalise_budget_exec(settings: Settings) -> BudgetExecMetrics:
    """Normalise local budget execution files into processed Parquet tables."""
    repo_root = settings.config_path.parent.parent if settings.config_path else find_repo_root()
    config = load_budget_columns_config(repo_root)

    processed_dir = settings.data_dirs()["processed"]
    raw_dir = settings.data_dirs()["raw"] / "budget_exec"

    uat_path = processed_dir / "uat.parquet"
    if not uat_path.is_file():
        raise FileNotFoundError(f"Missing required UAT reference table: {uat_path}")
    uat_ref = duckdb.query(
        f"SELECT siruta, name, county_name, county_code FROM '{uat_path}'"
    ).to_df()

    siruta_overrides = load_siruta_overrides(repo_root)

    # Locate manifest files in raw_dir
    manifest_entries_by_file: dict[str, dict[str, Any]] = {}
    for manifest_path in raw_dir.glob("**/manifest.json"):
        try:
            entries = json.loads(manifest_path.read_text(encoding="utf-8"))
            for entry in entries:
                manifest_entries_by_file[entry["filename"]] = entry
        except Exception as exc:
            logger.warning("Could not read manifest at %s: %s", manifest_path, exc)

    # Find raw Excel files
    excel_files = sorted(raw_dir.glob("**/*.xlsx")) + sorted(raw_dir.glob("**/*.xls"))
    if not excel_files:
        raise FileNotFoundError(f"No raw budget execution Excel files found under {raw_dir}")

    all_line_records: list[BudgetEuLineRecord] = []
    all_summary_records: list[UatEuReceivedYearRecord] = []

    years_processed: list[int] = []
    total_raw_rows_all = 0
    unmatched_details_all: list[dict[str, Any]] = []

    total_rec_by_year: dict[int, Decimal] = {}
    total_spent_by_year: dict[int, Decimal] = {}
    outliers_by_year: dict[int, int] = {}

    for file_path in excel_files:
        # Determine year from file name or config
        file_name = file_path.name
        year_match = None
        for yr_str in config:
            if yr_str in file_name:
                year_match = int(yr_str)
                break
        if year_match is None:
            # Default to 2024 if matching 'situaia-veniturilor-i-cheltuielilor-2024'
            if "2024" in file_name:
                year_match = 2024
            else:
                logger.info("Skipping unrecognized budget file: %s", file_name)
                continue

        year = year_match
        years_processed.append(year)
        year_cfg = config.get(str(year))
        if not year_cfg:
            raise ValueError(
                f"No column configuration found for year {year} in config/budget_exec_columns.yaml"
            )

        manifest_entry = manifest_entries_by_file.get(file_name, {})
        retrieved_at_val = (
            dt.datetime.fromisoformat(manifest_entry["retrieved_at"])
            if "retrieved_at" in manifest_entry
            else dt.datetime.now(dt.UTC)
        )
        provenance = {
            "source_id": "budget_exec",
            "source_url": manifest_entry.get(
                "url", "https://data.gov.ro/dataset/situatia-veniturilor-si-cheltuielilor-2024"
            ),
            "retrieved_at": retrieved_at_val,
            "raw_sha256": manifest_entry.get("sha256", ""),
        }

        logger.info("Parsing budget execution for year %d from %s", year, file_path)
        lines, summaries, info = parse_annual_budget_file(
            file_path=file_path,
            year=year,
            year_config=year_cfg,
            uat_ref=uat_ref,
            siruta_overrides=siruta_overrides,
            provenance=provenance,
        )

        all_line_records.extend(lines)
        all_summary_records.extend(summaries)
        total_raw_rows_all += info["total_raw_rows"]
        unmatched_details_all.extend(info["unmatched"])

        rec_sum = sum((s.eu_received_total for s in summaries), Decimal("0.00"))
        spent_sum = sum((s.eu_spent_total for s in summaries), Decimal("0.00"))
        total_rec_by_year[year] = rec_sum
        total_spent_by_year[year] = spent_sum

        # Flag ratio outliers (e.g. ratio > 10.0 or ratio < 0.1 where both > 0)
        outlier_count = sum(
            1
            for s in summaries
            if s.ratio_2014_2020 is not None
            and (s.ratio_2014_2020 > 10.0 or s.ratio_2014_2020 < 0.1)
        )
        outliers_by_year[year] = outlier_count

    # Write data/processed/uat_budget_eu_lines.parquet
    lines_schema = pa.schema(
        [
            ("siruta", pa.int64()),
            ("year", pa.int32()),
            ("classification_code", pa.string()),
            ("label", pa.string()),
            ("line_type", pa.string()),
            ("amount", pa.decimal128(20, 2)),
            ("unit", pa.string()),
            ("source_id", pa.string()),
            ("source_url", pa.string()),
            ("retrieved_at", pa.timestamp("us", tz="UTC")),
            ("raw_sha256", pa.string()),
        ]
    )

    lines_data = {
        "siruta": [r.siruta for r in all_line_records],
        "year": [r.year for r in all_line_records],
        "classification_code": [r.classification_code for r in all_line_records],
        "label": [r.label for r in all_line_records],
        "line_type": [r.line_type for r in all_line_records],
        "amount": [r.amount for r in all_line_records],
        "unit": [r.unit for r in all_line_records],
        "source_id": [r.source_id for r in all_line_records],
        "source_url": [r.source_url for r in all_line_records],
        "retrieved_at": [r.retrieved_at for r in all_line_records],
        "raw_sha256": [r.raw_sha256 for r in all_line_records],
    }
    lines_table = pa.Table.from_pydict(lines_data, schema=lines_schema)
    lines_output_path = processed_dir / "uat_budget_eu_lines.parquet"
    pq.write_table(lines_table, lines_output_path)
    logger.info("Wrote %d rows to %s", len(lines_table), lines_output_path)

    # Write data/processed/uat_eu_received_year.parquet
    summary_schema = pa.schema(
        [
            ("siruta", pa.int64()),
            ("year", pa.int32()),
            ("uat_name", pa.string()),
            ("county_name", pa.string()),
            ("county_code", pa.string()),
            ("eu_received_total", pa.decimal128(20, 2)),
            ("eu_received_2014_2020", pa.decimal128(20, 2)),
            ("eu_received_unspecified", pa.decimal128(20, 2)),
            ("eu_received_other", pa.decimal128(20, 2)),
            ("eu_spent_total", pa.decimal128(20, 2)),
            ("eu_spent_2014_2020", pa.decimal128(20, 2)),
            ("eu_spent_unspecified", pa.decimal128(20, 2)),
            ("eu_spent_pnrr_grant", pa.decimal128(20, 2)),
            ("eu_spent_pnrr_loan", pa.decimal128(20, 2)),
            ("spending_to_revenue_ratio", pa.float64()),
            ("ratio_2014_2020", pa.float64()),
            ("unit", pa.string()),
            ("source_id", pa.string()),
            ("source_url", pa.string()),
            ("retrieved_at", pa.timestamp("us", tz="UTC")),
            ("raw_sha256", pa.string()),
        ]
    )

    summary_data = {
        "siruta": [r.siruta for r in all_summary_records],
        "year": [r.year for r in all_summary_records],
        "uat_name": [r.uat_name for r in all_summary_records],
        "county_name": [r.county_name for r in all_summary_records],
        "county_code": [r.county_code for r in all_summary_records],
        "eu_received_total": [r.eu_received_total for r in all_summary_records],
        "eu_received_2014_2020": [r.eu_received_2014_2020 for r in all_summary_records],
        "eu_received_unspecified": [r.eu_received_unspecified for r in all_summary_records],
        "eu_received_other": [r.eu_received_other for r in all_summary_records],
        "eu_spent_total": [r.eu_spent_total for r in all_summary_records],
        "eu_spent_2014_2020": [r.eu_spent_2014_2020 for r in all_summary_records],
        "eu_spent_unspecified": [r.eu_spent_unspecified for r in all_summary_records],
        "eu_spent_pnrr_grant": [r.eu_spent_pnrr_grant for r in all_summary_records],
        "eu_spent_pnrr_loan": [r.eu_spent_pnrr_loan for r in all_summary_records],
        "spending_to_revenue_ratio": [r.spending_to_revenue_ratio for r in all_summary_records],
        "ratio_2014_2020": [r.ratio_2014_2020 for r in all_summary_records],
        "unit": [r.unit for r in all_summary_records],
        "source_id": [r.source_id for r in all_summary_records],
        "source_url": [r.source_url for r in all_summary_records],
        "retrieved_at": [r.retrieved_at for r in all_summary_records],
        "raw_sha256": [r.raw_sha256 for r in all_summary_records],
    }
    summary_table = pa.Table.from_pydict(summary_data, schema=summary_schema)
    summary_output_path = processed_dir / "uat_eu_received_year.parquet"
    pq.write_table(summary_table, summary_output_path)
    logger.info("Wrote %d rows to %s", len(summary_table), summary_output_path)

    matched_count = len(all_summary_records) - len(unmatched_details_all)

    return BudgetExecMetrics(
        years_processed=years_processed,
        total_raw_rows=total_raw_rows_all,
        total_uats_produced=len(all_summary_records),
        matched_uats_count=matched_count,
        unmatched_uats_count=len(unmatched_details_all),
        unmatched_details=unmatched_details_all,
        total_eu_received_by_year=total_rec_by_year,
        total_eu_spent_by_year=total_spent_by_year,
        outlier_count_by_year=outliers_by_year,
    )
